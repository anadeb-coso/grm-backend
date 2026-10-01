"""Exports de l'espace web : PDF mensuel (liste des CVD confirmés / non confirmés + attestation
signée), PDF par CVD, et Excel mensuel."""
import io

from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa

from . import services
from .models import render_message


class PdfGenerationError(Exception):
    pass


def _render_pdf(template_name, context):
    html = render_to_string(template_name, context)
    buffer = io.BytesIO()
    result = pisa.CreatePDF(html, dest=buffer, encoding='utf-8')
    if result.err:
        raise PdfGenerationError(f'{result.err} erreur(s) lors de la génération du PDF')
    return buffer.getvalue()


def _signature_block(options):
    """`options` : dict validé par dashboard/internet_credits/forms.py::ExportOptionsForm."""
    drawn = options.get('drawn_signature') if options.get('signature_mode') == 'drawn' else ''
    return {
        'place': options.get('place') or '',
        'date': timezone.localdate(),
        'signatory_name': options.get('signatory_name') or '',
        'signatory_title': options.get('signatory_title') or '',
        'drawn_signature': drawn,
    }


def monthly_pdf(campaign, options):
    lines = services.campaign_lines(campaign)
    stats = services.campaign_stats(lines)
    context = campaign.message_context(total=stats['total'], confirmes=stats['cvgp_confirmed'])
    return _render_pdf('internet_credits/pdf/monthly.html', {
        'campaign': campaign,
        'confirmed': [line for line in lines if line['cvgp']],
        'pending': [line for line in lines if not line['cvgp']],
        'stats': stats,
        'attestation': render_message(options.get('text') or campaign.attestation_text, context),
        'signature': _signature_block(options),
        'generated_at': timezone.localtime(),
    })


def cvd_pdf(beneficiary, options):
    campaign = beneficiary.campaign
    cvgp = services.cvgp_confirmation(beneficiary)
    context = campaign.message_context(
        cvd=beneficiary.label, villages=beneficiary.village_names,
        nom=beneficiary.member_name, telephone=beneficiary.phone_number,
    )
    return _render_pdf('internet_credits/pdf/cvd.html', {
        'campaign': campaign,
        'beneficiary': beneficiary,
        'cvgp': cvgp,
        'fcs': services.fc_confirmations(beneficiary),
        'agreement': render_message(options.get('text') or campaign.agreement_text, context),
        'signature': _signature_block(options),
        'generated_at': timezone.localtime(),
    })


def _fmt(dt):
    return timezone.localtime(dt).strftime('%d/%m/%Y %H:%M') if dt else ''


def monthly_excel(campaign):
    lines = services.campaign_lines(campaign)
    stats = services.campaign_stats(lines)

    wb = Workbook()
    ws = wb.active
    ws.title = campaign.month_label[:31]

    ws.append([f'Confirmation de réception du forfait internet CVGP — {campaign.month_label}'])
    ws['A1'].font = Font(bold=True, size=14, color='1A9C6E')
    ws.append([
        f"Date d'envoi : {campaign.sent_date.strftime('%d/%m/%Y')}",
        f'Forfait : {campaign.package_label}',
        f'Montant : {campaign.amount_display}',
        f"CVGP confirmés : {stats['cvgp_confirmed']}/{stats['total']}",
        f"FC/AC confirmés : {stats['fc_confirmed']}/{stats['total']}",
    ])
    ws.append([])

    headers = [
        'N°', 'CVD', 'Villages', 'Membre CVGP', 'Téléphone',
        'Statut CVGP', 'Date confirmation CVGP', 'Message coché (CVGP)', 'Description (CVGP)', 'Signature CVGP',
        'FC/AC', 'Statut FC/AC', 'Date confirmation FC/AC', 'Message coché (FC/AC)', 'Description (FC/AC)',
        'Signature FC/AC',
    ]
    ws.append(headers)
    header_row = ws.max_row
    fill = PatternFill('solid', fgColor='E8F6F0')
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')

    for line in lines:
        b, cvgp, fc = line['beneficiary'], line['cvgp'], line['fc']
        ws.append([
            line['index'], b.label, b.village_names, b.member_name, b.phone_number,
            'Confirmé' if cvgp else 'Non confirmé', _fmt(cvgp.confirmed_at) if cvgp else '',
            cvgp.message if cvgp else '', (cvgp.description or '') if cvgp else '',
            'Oui' if cvgp and cvgp.signature else 'Non',
            ', '.join(c.confirmed_by_name for c in line['fcs']),
            'Confirmé' if fc else 'Non confirmé', _fmt(fc.confirmed_at) if fc else '',
            fc.message if fc else '', (fc.description or '') if fc else '',
            'Oui' if fc and fc.signature else 'Non',
        ])

    widths = [5, 22, 30, 24, 15, 13, 18, 50, 30, 10, 24, 13, 18, 50, 30, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
