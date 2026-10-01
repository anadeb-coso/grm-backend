import base64
import io
import os

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from internet_credits import exports, services
from internet_credits.models import InternetCreditBeneficiary

from .conftest import confirmation_record

pytestmark = pytest.mark.django_db


def _signature_png(width=600, height=220):
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    draw.line([(20, 150), (120, 60), (220, 170), (340, 50), (460, 160), (580, 80)], fill='black', width=6)
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(buffer.getvalue()).decode()


@pytest.fixture
def confirmed_campaign(actors, campaign):
    line = InternetCreditBeneficiary.objects.get(campaign=campaign, cvgp_member=actors['cvgp'])
    signature = _signature_png()
    services.create_confirmation(actors['cvgp'], confirmation_record(line, signature=signature, description='Reçu le 3'), 'grm_mobile')
    services.create_confirmation(actors['fc'], confirmation_record(line, signature=signature), 'cdd_mobile')
    return campaign, line


def _dump(name, content):
    """`IC_EXPORT_DUMP_DIR=<dossier>` pour inspecter visuellement les fichiers générés."""
    target = os.environ.get('IC_EXPORT_DUMP_DIR')
    if target:
        with open(os.path.join(target, name), 'wb') as f:
            f.write(content)


def test_monthly_pdf_with_drawn_signature(confirmed_campaign):
    campaign, _ = confirmed_campaign
    pdf = exports.monthly_pdf(campaign, {
        'text': 'Je confirme que ces CVD ont reçu le forfait de {mois} {annee}.',
        'place': 'Lomé', 'signatory_name': 'M. TINDAME', 'signatory_title': 'Coordonnateur',
        'signature_mode': 'drawn', 'drawn_signature': _signature_png(),
    })
    _dump('monthly.pdf', pdf)
    assert pdf.startswith(b'%PDF') and b'/Image' in pdf


def test_monthly_pdf_with_blank_signature_box(campaign):
    pdf = exports.monthly_pdf(campaign, {'signature_mode': 'blank'})
    _dump('monthly_blank.pdf', pdf)
    assert pdf.startswith(b'%PDF')


def test_cvd_pdf(confirmed_campaign):
    _, line = confirmed_campaign
    pdf = exports.cvd_pdf(line, {'signature_mode': 'blank', 'signatory_name': 'M. LARE'})
    _dump('cvd.pdf', pdf)
    assert pdf.startswith(b'%PDF') and b'/Image' in pdf


def test_monthly_excel(confirmed_campaign):
    campaign, _ = confirmed_campaign
    content = exports.monthly_excel(campaign)
    _dump('monthly.xlsx', content)
    ws = load_workbook(io.BytesIO(content)).active
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    by_cvd = {row[1]: row for row in rows}
    assert by_cvd['KPEKPE'][5] == 'Confirmé' and by_cvd['KPEKPE'][9] == 'Oui' and by_cvd['KPEKPE'][11] == 'Confirmé'
    assert by_cvd['DAPAONG'][5] == 'Non confirmé'
